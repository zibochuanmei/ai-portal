from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "AI_Portal项目骨架阶段四人同步开发分工与整合计划_V1.0.docx"
BLACK = "000000"
BLUE = "1F4E78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F5F7FA"
BORDER = "D9E2F3"


def set_run_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None, color: str = BLACK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, *, header: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.05
    run = paragraph.add_run(text)
    set_run_font(run, size=9.2 if header else 9.3, bold=header, color="FFFFFF" if header else BLACK)
    if header:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_borders(cell)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for index, header in enumerate(headers):
        set_cell_shading(table.rows[0].cells[index], BLUE)
        set_cell_text(table.rows[0].cells[index], header, header=True)
        if widths:
            table.rows[0].cells[index].width = Inches(widths[index])
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            set_cell_shading(cells[index], LIGHT_BLUE if row_index % 2 == 0 else "FFFFFF")
            set_cell_text(cells[index], value)
            if widths:
                cells[index].width = Inches(widths[index])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color=BLACK)


def add_body(doc: Document, text: str, *, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        prefix = paragraph.add_run(bold_prefix)
        set_run_font(prefix, size=10.5, bold=True)
        body = paragraph.add_run(text[len(bold_prefix):])
        set_run_font(body, size=10.5)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=10.5)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.05
        run = paragraph.add_run(item)
        set_run_font(run, size=10.2)


def add_footer(section) -> None:
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("AI Portal | 项目骨架阶段四人同步开发分工与整合计划")
    set_run_font(run, size=8.5, color="666666")


def build_document() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)
    add_footer(section)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("AI Portal 项目骨架阶段四人同步开发分工与整合计划")
    set_run_font(run, size=20, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("基础平台阶段，不包含 RAG 与 Agent 业务设计")
    set_run_font(run, size=11, color="4F4F4F")

    meta = doc.add_table(rows=2, cols=4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = True
    values = [["版本", "V1.0", "编制日期", "2026 年 9 月 21 日"], ["适用团队", "四人研发团队", "协作方式", "两个月同步推进，之后共同设计 Agent"]]
    for r_idx, row in enumerate(values):
        for c_idx, value in enumerate(row):
            cell = meta.rows[r_idx].cells[c_idx]
            set_cell_shading(cell, LIGHT_GRAY if c_idx % 2 == 0 else "FFFFFF")
            set_cell_text(cell, value)
            if c_idx % 2 == 0:
                for run in cell.paragraphs[0].runs:
                    set_run_font(run, size=9.2, bold=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

    add_heading(doc, "1 文档目的与阶段边界")
    add_body(doc, "本计划用于指导四人团队在 AI Portal 项目骨架阶段同步开发。前两个月由两组并行推进后端基础、前端页面、权限、部署和测试，避免等待式开发；两个月结束后先完成基础平台验收，再由四人共同进入具体 Agent 的需求和技术设计阶段。")
    add_body(doc, "前两个月建设范围包括项目目录、前后端运行链路、用户与角色、部门和 Agent 授权基础、文件与任务接口占位、日志、Docker 部署和测试。RAG 文档解析、向量检索、Supervisor、Prompt、Tool 和具体 Agent 业务逻辑在前两个月暂不纳入开发范围。")

    add_heading(doc, "2 协作原则")
    add_bullets(doc, [
        "两组同步开发：每周按照同一目标拆分任务，前后端使用约定好的 DTO 和接口并行实现。",
        "一人主做、一人复核：每项代码必须由本组另一名成员评审；公共接口由另一组至少一人参与评审。",
        "先定契约再写实现：接口、字段、状态和错误码写入 packages/contracts，变更必须同步更新文档。",
        "后端决定权限：前端只展示后端返回的结果，不能在前端自行判断员工是否有 Agent 权限。",
        "共同承担上线：两组都参与联调、权限测试、Docker 验证、部署和故障排查，不形成单点依赖。",
    ])

    add_heading(doc, "3 小组一 刘志博与李明山")
    add_body(doc, "小组一负责后端基础、数据结构、权限服务和运行环境，为小组二提供稳定的接口、测试账号和联调环境。当前只建设基础接口和权限框架，不实现 RAG 或 Agent 业务流程。")
    add_table(doc, ["工作方向", "主要交付物", "阶段验收"], [
        ["后端基础", "FastAPI 分层、统一配置、异常处理、REST/OpenAPI 接口、请求 trace_id", "后端可启动，/health、/me 等接口可调用，错误格式统一"],
        ["数据库基础", "SQLAlchemy/Alembic 结构；users、departments、groups、agent_grants、audit_logs 等表设计", "迁移可重复执行，初始化数据不重复，字段和索引通过评审"],
        ["认证与权限", "employee/platform_admin 两种角色；用户停用；Agent 有效授权查询和执行校验", "停用用户无法访问，未授权 Agent 列表和执行请求均被拒绝"],
        ["文件与任务接口", "文件上传占位、任务状态模型、结果查询和 SSE 接口约定", "接口能被前端联调，状态流转和错误码明确"],
        ["部署与日志", "Docker Compose、Nginx、后端 JSONL 日志、日志目录挂载、环境变量模板", "本地和服务器可重复启动，能用 trace_id 定位请求"],
        ["后端测试", "健康检查、权限、停用账号、越权访问和接口参数测试", "合并前后端基础测试通过，失败原因可从日志定位"],
    ], widths=[1.25, 3.0, 2.5])
    add_heading(doc, "3.1 小组内协作重点", level=2)
    add_bullets(doc, [
        "李明山优先推进 API、数据模型、用户组织和 PermissionService；刘志博同步推进 Docker、日志、环境配置、测试工具和部署脚本。",
        "数据库迁移、权限规则和公共 DTO 不允许单人直接合并，必须由两人共同评审。",
        "每周向小组二提供可运行的接口版本、演示账号、错误码和变更说明。",
    ])

    add_heading(doc, "4 小组二 符润与熊赵妤")
    add_body(doc, "小组二负责 Vue 员工端、管理员端、接口封装和前端联调体验。页面先使用当前演示接口和模拟数据开发，后端接口稳定后切换到真实权限结果。")
    add_table(doc, ["工作方向", "主要交付物", "阶段验收"], [
        ["前端基础", "Vue 3、TypeScript、Vite、Router、Pinia、Element Plus、统一布局", "前端可启动，路由和状态结构清晰，常用屏幕无明显布局问题"],
        ["登录与员工端", "登录页、当前用户状态、员工工作台、Agent 列表和文件上传交互", "员工能看到自己的身份和有效 Agent，未授权入口有明确提示"],
        ["管理员端", "用户、部门、账号启停、Agent 授权、授权预览、审计记录页面", "管理员页面按角色展示，普通员工不能进入管理功能"],
        ["接口封装", "Axios 客户端、统一错误提示、Loading/Empty/Error 状态、trace_id 展示和问题反馈", "接口异常可提示，前端能记录 request_id 供排查"],
        ["任务与文件体验", "上传状态、任务处理中、成功、失败、取消和历史记录页面骨架", "刷新页面后能根据任务接口恢复基本状态"],
        ["前端测试", "组件测试、路由权限测试和基础 Playwright 流程", "登录、员工端、管理员端主路径可重复验证"],
    ], widths=[1.25, 3.0, 2.5])
    add_heading(doc, "4.1 小组内协作重点", level=2)
    add_bullets(doc, [
        "符润优先推进员工端、登录状态、API 客户端和 Pinia；熊赵妤同步推进管理员端、权限展示、异常状态和前端测试。",
        "页面不能自行合并授权规则；所有可见 Agent、管理员操作和账号状态以 API 返回为准。",
        "每周向小组一提供页面录屏或可操作环境、接口调用清单和待解决的交互问题。",
    ])

    add_heading(doc, "5 同步开发节奏")
    add_table(doc, ["阶段", "小组一 后端与环境", "小组二 前端与联调", "共同整合点"], [
        ["第 1-2 周 项目基线", "目录、配置、日志、健康检查、Compose、接口 DTO", "页面布局、路由、Pinia、API 客户端、登录页骨架", "确认分支、接口文档、环境变量和演示账号"],
        ["第 3-4 周 身份与权限", "用户/部门/角色/授权模型和后端校验", "登录状态、角色路由、Agent 列表和权限提示", "员工与管理员两条路径完成首次联调"],
        ["第 5-6 周 文件与任务", "文件、任务、事件、结果接口和审计记录", "上传、处理中、成功/失败、历史和错误状态", "完成一次文件任务端到端演示"],
        ["第 7-8 周 测试与交付", "越权测试、日志排查、备份、Docker 和服务器部署", "响应式优化、异常场景、Playwright 流程", "完成基础平台验收、部署演练和问题清单闭环"],
        ["两个月后 Agent 设计", "参与 AgentExecutionService、权限上下文和运行接口设计", "参与 Agent 工作区、输入输出和状态交互设计", "四人共同确定 Agent 范围、技术方案和后续排期"],
    ], widths=[1.35, 2.15, 2.15, 2.1])
    add_body(doc, "同步规则：前两个月每周至少一次接口联调、一次越权测试和一次部署验证。联调问题在当天登记到项目看板，涉及公共 DTO、权限规则或数据库迁移的变更必须由两组共同确认。两个月结束时冻结基础接口，经过验收后再启动 Agent 设计。")

    add_heading(doc, "6 基础接口与代码整合规则")
    add_bullets(doc, [
        "前端只调用 Portal API，不直接访问 PostgreSQL、Redis、MinIO、Milvus 或模型服务。",
        "后端路由只负责参数校验和响应编排，业务逻辑进入 Service，数据库访问进入 Repository。",
        "Agent 列表、任务创建、文件下载和管理员操作都必须在后端重新校验用户身份与权限。",
        "分支建议使用 codex/feature-xxx 或 feature/xxx；每个 PR 至少由本组一人和另一组一人评审。",
        "合并前必须通过 Python 编译、后端测试、前端构建、权限测试和 Compose 配置检查。",
    ])
    add_table(doc, ["公共文件或模块", "维护方式"], [
        ["packages/contracts/api-v1.md", "接口路径、请求响应、错误码和状态变更由两组共同维护"],
        ["backend/app/core、permission、数据库迁移", "小组一主维护，小组二参与评审"],
        ["frontend/src/api、router、stores", "小组二主维护，小组一参与接口和权限评审"],
        ["infra、Docker、日志和部署文档", "小组一主维护，小组二参与启动验证"],
        ["tests/与端到端流程", "两组共同维护，测试失败必须由对应功能小组修复"],
    ], widths=[2.6, 4.2])

    add_heading(doc, "7 两个月后 Agent 模块设计与整合准备")
    add_body(doc, "前两个月不直接开发具体 Agent。基础平台通过验收后，四人共同进入 Agent 设计阶段，先确定业务场景、输入输出、权限范围、模型调用方式、工作流边界和验收样例，再进入实现排期。")
    add_bullets(doc, [
        "设计阶段首先确定首个 Agent 的业务目标、用户、文件类型、输出格式和成功标准。",
        "Agent 模块必须通过统一执行服务启动，不能从前端或单个路由直接调用模型。",
        "Agent 的输入、输出、文件和错误状态沿用前两个月冻结的 DTO、权限、日志和任务链路。",
        "后续接入 LangGraph、RAG 或 Supervisor 时，继续复用当前权限、日志、任务和审计基础。",
        "每增加一个 Agent，必须同时提交接口说明、权限范围、测试样例、日志字段和回滚方案。",
    ])

    add_heading(doc, "8 基础阶段完成标准")
    add_bullets(doc, [
        "员工和管理员可以登录并进入各自页面；停用账号无法继续访问。",
        "管理员可以维护用户状态、部门和 Agent 授权；普通员工只能看到后端授权结果。",
        "前端可以完成文件上传、任务提交、状态展示和结果查看的基础流程。",
        "后端请求、权限拒绝和任务异常都有 trace_id，可以在日志文件中定位。",
        "项目可以通过 Docker Compose 启动，四人电脑和测试服务器使用同一套环境配置方式。",
        "代码、接口、数据库迁移和部署文档均已提交 Git，并通过两组交叉评审。",
        "前两个月基础平台验收通过后，四人共同输出 Agent 设计方案和下一阶段排期。",
    ])

    doc.core_properties.title = "AI Portal 项目骨架阶段四人同步开发分工与整合计划"
    doc.core_properties.subject = "AI Portal 基础平台阶段团队分工"
    doc.core_properties.author = "AI Portal 项目组"
    doc.core_properties.comments = "前两个月两组同步推进，基础平台验收后共同设计 Agent 模块。"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
